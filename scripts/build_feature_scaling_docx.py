#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build feature-scaling experiment Word report."""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "data" / "figures"
OUT = ROOT / "data" / "特征扩增实验报告.docx"
BLUE = RGBColor(0x2E,0x74,0xB5)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = Inches(w/1440)


def add_heading(doc, text, level=1):
    h = doc.add_heading(level=level); r = h.add_run(text)
    r.font.name="Calibri"; r.font.size=Pt(16 if level==1 else 13); r.font.bold=True; r.font.color.rgb=BLUE
    h.paragraph_format.space_before=Pt(12); h.paragraph_format.space_after=Pt(6)


def add_body(doc, text):
    p=doc.add_paragraph(); r=p.add_run(text); r.font.name="Calibri"; r.font.size=Pt(11)
    p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.2


def add_figure(doc, path, caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=p.add_run(); run.add_picture(str(path), width=Inches(6.0))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(caption); r.font.name="Calibri"; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x55,0x55,0x55)


def main():
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11)

    t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=t.add_run("特征扩增实验报告"); r.font.name="Calibri"; r.font.size=Pt(18); r.font.bold=True

    add_heading(doc,"1. 特征提取算法说明",1)
    add_body(doc,"低方差过滤：去除几乎不变化的特征。F 值筛选：使用 ANOVA F 统计量，衡量特征与类别之间的线性相关性，选择 F 值最大的 k 个特征。L1 正则化：使用带 L1 惩罚的线性 SVM，通过稀疏权重自动选择特征。NSRE：新对称相对熵，度量正负类样本特征分布的距离，选择分布差异大的特征。组合方法：FClassif_NSRE 先粗筛再精筛；L1_NSRE 先 L1 后 NSRE。")

    add_heading(doc,"2. 机器学习与 MLP 算法说明",1)
    add_body(doc,"LogisticRegression：线性分类器，使用 L2 正则，最大迭代 2000。RandomForest：集成决策树。GradientBoosting：梯度提升树。SVC：支持向量机。KNN：k 近邻。MLP：多层感知器，结构为输入层-128-ReLU-Dropout-64-ReLU-Dropout-输出层，Adam 优化，学习率 0.001，权重衰减 1e-4。")

    add_heading(doc,"3. 实验结果",1)
    add_body(doc,"结果采用 5 折分层交叉验证，随机种子 42。")
    add_figure(doc, FIG/"feature_scaling_accuracy.png", "图 1. FClassif_NSRE + LogisticRegression 的 Accuracy（mean ± std）")
    add_figure(doc, FIG/"feature_scaling_macro_f1.png", "图 2. FClassif_NSRE + LogisticRegression 的 Macro-F1（mean ± std）")

    doc.save(OUT); print(f"Saved -> {OUT}")


if __name__=="__main__":
    main()

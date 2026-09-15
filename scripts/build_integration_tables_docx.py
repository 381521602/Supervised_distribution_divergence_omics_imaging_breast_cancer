#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build integration results tables Word document."""

from pathlib import Path
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/"data"/"多组学整合结果表.docx"
BLUE=RGBColor(0x2E,0x74,0xB5)


def set_cell_shading(cell, fill):
    tc_pr=cell._tc.get_or_add_tcPr(); shd=tc_pr.find(qn("w:shd"))
    if shd is None:
        shd=OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_widths(table, widths):
    table.autofit=False
    for row in table.rows:
        for idx,w in enumerate(widths):
            row.cells[idx].width=Inches(w/1440)


def add_heading(doc,text):
    h=doc.add_heading(level=1); r=h.add_run(text)
    r.font.name="Calibri"; r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=BLUE


def add_table(doc, df, group_col, models, metrics):
    table=doc.add_table(rows=1, cols=2+len(metrics))
    hdr=table.rows[0].cells
    hdr[0].text=group_col; hdr[1].text="模型"
    for i,m in enumerate(metrics, start=2): hdr[i].text=m
    for cell in hdr: set_cell_shading(cell,"F2F4F7")
    for group in sorted(df[group_col].unique()):
        for model in models:
            sub=df[(df[group_col]==group)&(df.model==model)]
            if sub.empty: continue
            cells=table.add_row().cells
            cells[0].text=group; cells[1].text=model
            for i,m in enumerate(metrics, start=2):
                r=sub[sub.metric==m]
                if len(r)>0:
                    row=r.iloc[0]; cells[i].text=f"{row['mean']:.4f} ± {row['std']:.4f}"
    set_table_widths(table,[2000]+[2000 for _ in range(len(metrics))]+[1800])


def main():
    trad=pd.read_csv(ROOT/"data/integration_final_datasets_results.tsv",sep="\t")
    mlp=pd.read_csv(ROOT/"data/integration_mlp_cox_results.tsv",sep="\t")
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11)
    title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=title.add_run("多组学整合结果表"); r.font.name="Calibri"; r.font.size=Pt(17); r.font.bold=True

    add_heading(doc,"传统机器学习整合")
    add_heading(doc,"PAM50 分型")
    add_table(doc, trad[trad.task=="PAM50"], "combo", ["LogisticRegression","RandomForest","GradientBoosting","SVC","KNN"], ["accuracy","macro_f1"])
    add_heading(doc,"生存预测")
    add_table(doc, trad[trad.task=="Survival"], "combo", ["LogisticRegression","RandomForest","GradientBoosting","SVC","KNN"], ["roc_auc","c_index"])

    add_heading(doc,"MLP 与 CoxPH 整合")
    add_heading(doc,"PAM50 分型")
    add_table(doc, mlp[mlp.task=="PAM50"], "combo", ["MLP"], ["accuracy","macro_f1"])
    add_heading(doc,"生存预测")
    add_table(doc, mlp[mlp.task=="Survival"], "combo", ["CoxPH"], ["roc_auc","c_index"])

    doc.save(OUT); print(f"Saved -> {OUT}")


if __name__=="__main__":
    main()

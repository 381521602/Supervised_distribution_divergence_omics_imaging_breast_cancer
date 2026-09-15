#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build imaging method and examples report."""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT=Path(__file__).resolve().parent.parent
IMG=ROOT/"data/images/examples"
OUT=ROOT/"data"/"组学图像化方法与示例.docx"
BLUE=RGBColor(0x2E,0x74,0xB5)


def set_cell_shading(cell,fill):
    tc_pr=cell._tc.get_or_add_tcPr(); shd=tc_pr.find(qn("w:shd"))
    if shd is None:
        shd=OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"),fill)


def set_table_widths(table,widths):
    table.autofit=False
    for row in table.rows:
        for idx,w in enumerate(widths):
            row.cells[idx].width=Inches(w/1440)


def add_heading(doc,text,level=1):
    h=doc.add_heading(level=level); r=h.add_run(text)
    r.font.name="Calibri"; r.font.size=Pt(16 if level==1 else 13); r.font.bold=True; r.font.color.rgb=BLUE
    h.paragraph_format.space_before=Pt(12); h.paragraph_format.space_after=Pt(6)


def add_body(doc,text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; r=p.add_run(text); r.font.name="Calibri"; r.font.size=Pt(11)
    p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.2


def add_figure(doc,path,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=p.add_run(); run.add_picture(str(path),width=Inches(6.0))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(caption); r.font.name="Calibri"; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x55,0x55,0x55)


def main():
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    normal=doc.styles["Normal"]; normal.font.name="Calibri"; normal.font.size=Pt(11)
    title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=title.add_run("组学数据图像化方法与示例"); r.font.name="Calibri"; r.font.size=Pt(17); r.font.bold=True

    add_heading(doc,"1. 图像化方法",1)
    add_body(doc,"图像化方法的目标是将高维一维组学特征转换为二维灰度图，以便卷积神经网络提取局部和全局结构信息。具体步骤如下：")
    add_body(doc,"第一步，对每个特征计算 NSRE 分数。NSRE 通过比较正负类样本特征值概率分布的差异来衡量该特征的判别能力，分数越高表示该特征在不同类别间的分布差异越大。")
    add_body(doc,"第二步，按照 NSRE 分数从高到低对所有特征进行排序。排序越靠前的特征，判别能力越强。")
    add_body(doc,"第三步，采用中心向外的螺旋顺序将排序后的特征逐个填入方形网格。NSRE 分数最高的特征被放置在图像中心，随着排序位置下降，特征逐渐向外围扩展。这种排列方式使最重要的特征集中在图像中心附近，有利于卷积运算重点关注核心区域，同时降低边界效应的影响。")
    add_body(doc,"第四步，将标准化后的特征值直接作为像素灰度值；如果特征数量少于网格容量，则剩余位置补 0。")
    add_figure(doc, IMG/"imaging_method_diagram.png", "图 1. 组学数据图像化流程")
    add_body(doc,"图像尺寸根据最终特征数确定：PAM50 中 mRNA 为 20×20、CNV 为 8×8、miRNA 为 25×25；生存预测中 mRNA 为 15×15、CNV 为 13×13、miRNA 为 25×25。")

    add_heading(doc,"2. PAM50 亚型示例",1)
    for omics in ["mRNA","CNV","miRNA"]:
        add_figure(doc, IMG/f"PAM50_{omics}_by_class.png", f"PAM50 {omics} 按亚型的灰度图示例")
        add_figure(doc, IMG/f"PAM50_{omics}_class_average.png", f"PAM50 {omics} 各亚型平均灰度图")

    add_heading(doc,"3. 生存事件示例",1)
    for omics in ["mRNA","CNV","miRNA"]:
        add_figure(doc, IMG/f"Survival_{omics}_by_class.png", f"Survival {omics} 按生存事件的灰度图示例")
        add_figure(doc, IMG/f"Survival_{omics}_class_average.png", f"Survival {omics} 各事件平均灰度图")

    doc.save(OUT); print(f"Saved -> {OUT}")


if __name__=="__main__":
    main()

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a draft research article DOCX."""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "data" / "figures"
OUT = ROOT / "data" / "组学数据图像化_文章草稿.docx"
BLUE = RGBColor(0x2E, 0x74, 0xB5)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = Inches(w / 1440)


def add_heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    r.font.name = "Calibri"; r.font.size = Pt(16 if level == 1 else 13); r.font.bold = True; r.font.color.rgb = BLUE
    h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
    return h


def add_body(doc, text):
    p = doc.add_paragraph(); r = p.add_run(text); r.font.name = "Calibri"; r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.15
    return p


def add_figure(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(); run.add_picture(str(path), width=Inches(6.0))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.font.name = "Calibri"; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55,0x55,0x55)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("基于组学数据图像化与门控全尺寸卷积网络的乳腺癌多组学融合分析"); r.font.name = "Calibri"; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = RGBColor(0x0B,0x25,0x45)

    add_heading(doc, "摘要", 1)
    add_body(doc, "本研究提出一种基于组学数据图像化的乳腺癌多组学融合分析框架。首先对 TCGA-BRCA 的 mRNA、CNV、miRNA 数据进行样本对齐和特征筛选，采用新对称相对熵（JSD）进行特征排序与加权，并将一维特征转换为二维灰度图；随后构建全尺寸卷积多流网络，并引入门控机制和 Mish 激活。结果表明，PAM50 分子分型中门控 JSD 加权全尺寸卷积网络达到约 0.86 的准确率，生存预测中 mRNA 扩增特征后 Cox C-index 可达 0.64。关键基因挖掘识别出若干与乳腺癌相关的候选标志物，为图像化多组学融合提供了一种可解释的新方法。")

    add_heading(doc, "1. 引言", 1)
    add_body(doc, "乳腺癌具有高度分子异质性，多组学整合分析可揭示不同层次的分子机制。传统机器学习以向量形式输入高维组学特征，难以利用特征间的空间关系。本研究将组学特征图像化，并使用全尺寸卷积网络进行融合预测，同时引入 JSD 信息熵特征排序和门控机制，提升分类与生存预测性能。")

    add_heading(doc, "2. 数据与方法", 1)
    add_body(doc, "数据来源于 TCGA-BRCA 的 GDC 与 UCSC Xena。共 1098 例乳腺癌，其中 1073 例同时具备 mRNA、miRNA 和基因级 CNV。特征筛选采用方差过滤、F 值、L1、JSD 及其组合。图像化使用 JSD 排序与加权，将特征填入方形灰度图；模型采用多流全尺寸卷积，并加入门控机制与 Mish 激活。全部实验使用 5 折分层交叉验证，随机种子为 42。")

    add_heading(doc, "3. 结果", 1)
    add_heading(doc, "3.1 单组学预测", 2)
    add_body(doc, "PAM50 分型中，mRNA 单组学最优 Accuracy 约 0.86，CNV 和 miRNA 较弱。生存预测中 CNV 单组学 Cox C-index 最高，约 0.66。")
    add_figure(doc, FIG / "summary_pam50_accuracy.png", "图 1. PAM50 四分类单组学与多方法对比（含标准差）")
    add_figure(doc, FIG / "summary_survival.png", "图 2. 生存预测单组学与多方法对比（含标准差）")

    add_heading(doc, "3.2 图像化与网络结构消融", 2)
    add_body(doc, "消融实验显示，全尺寸卷积优于 3×3 残差、多尺度、Inception、MobileNet、Grouped Conv 等结构；JSD 像素加权将 PAM50 准确率从 0.813 提升至 0.860；门控机制进一步提升 Macro-F1 至 0.851。Mish 激活略优于 ReLU。")
    add_figure(doc, FIG / "fig4_adaptive_pipeline.png", "图 3. 自适应 JSD 特征筛选与图像化流程")

    add_heading(doc, "3.3 多组学整合", 2)
    add_body(doc, "生存预测中模型级晚期融合软投票 AUC 达 0.646，优于特征级早期融合；PAM50 分型中早期融合与 mRNA 单组学接近。")

    add_heading(doc, "4. 讨论与结论", 1)
    add_body(doc, "图像化深度学习框架在 PAM50 分型和 mRNA 生存预测上表现出竞争力，并具有较好的可解释性。关键基因挖掘识别出多个候选标志物。后续将进一步进行通路富集和外部验证。")

    doc.save(OUT)
    print(f"Saved -> {OUT}")


if __name__ == "__main__":
    main()

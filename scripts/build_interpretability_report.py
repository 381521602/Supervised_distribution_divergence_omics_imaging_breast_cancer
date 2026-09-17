#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build interpretability and key-factor mining DOCX report."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data/interpretability"
FIG = DATA / "figures"
OUT = ROOT / "可解释性与关键因子挖掘报告.docx"


def set_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5), 18, 10),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5), 14, 7),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
    p.paragraph_format.space_after = Pt(10)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(8)


def add_desc(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p.paragraph_format.space_after = Pt(8)


def set_table_width(table, widths):
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    for el in tblPr.findall(qn("w:tblW")):
        tblPr.remove(el)
    tblW = tblPr.makeelement(qn("w:tblW"), {qn("w:w"): str(int(sum(widths) * 1440)), qn("w:type"): "dxa"})
    tblPr.append(tblW)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths):
                cell.width = Inches(widths[i])


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    if widths is None:
        widths = [6.5 / len(headers)] * len(headers)
    set_table_width(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(9)
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)


def top_rows(df, cols, n=10):
    return [[round(float(x), 4) if isinstance(x, (float, int)) else str(x) for x in row] for row in df.head(n).itertuples(index=False)]


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1); s.bottom_margin = Inches(1); s.left_margin = Inches(1); s.right_margin = Inches(1)
    set_style(doc)
    add_title(doc, "可解释分析与关键因子挖掘报告")
    doc.add_paragraph("本报告基于 mRNA PAM50 和 Survival 模型，使用 SHAP、CNN 输入梯度重要性、跨折共识排序、单因素 Cox 和通路富集，识别与乳腺癌分子分型和预后相关的关键基因。")

    doc.add_heading("1. 方法", level=1)
    doc.add_paragraph("SHAP 用于解释 LogisticRegression；FullSizeCNN 由于全尺寸卷积输出为 1×1，使用输入梯度 saliency 映射像素到基因。所有分析均在 5 折交叉验证框架下汇总。Survival 关键因子使用线性替代模型系数和单因素 Cox。")

    doc.add_heading("2. SHAP 重要性", level=1)
    shap = pd.read_csv(DATA / "mrna_pam50_shap_importance.tsv", sep="\t")
    add_table(doc, ["Gene", "Mean importance", "Std importance", "Fold frequency"], top_rows(shap, shap.columns, 10), widths=[1.5, 1.6, 1.5, 1.4])
    doc.add_picture(str(FIG / "fig_shap_top.png"), width=Inches(6.2))
    add_caption(doc, "图 1. mRNA PAM50 LogisticRegression SHAP top-20 基因")
    add_desc(doc, "SHAP 重要性最高的基因包括 CYP2B7P1、TFF1、C1orf64、AGR3、KCNJ3、MIA、PPP1R14C、ESR1 等，均在 5 折中稳定进入 top-100。")

    doc.add_heading("3. CNN 像素重要性", level=1)
    cnn = pd.read_csv(DATA / "mrna_pam50_cnn_saliency_importance.tsv", sep="\t")
    add_table(doc, ["Gene", "Mean importance", "Std importance", "Fold frequency"], top_rows(cnn, cnn.columns, 10), widths=[1.5, 1.6, 1.5, 1.4])
    doc.add_picture(str(FIG / "fig_cnn_top.png"), width=Inches(6.2))
    add_caption(doc, "图 2. mRNA PAM50 CNN saliency top-20 基因")
    add_desc(doc, "CNN 输入梯度重要性反映了图像化模型中不同基因像素对 PAM50 分类的贡献。")

    doc.add_heading("4. 共识关键因子", level=1)
    cons = pd.read_csv(DATA / "mrna_pam50_consensus_key_factors.tsv", sep="\t")
    add_table(doc, ["Gene", "SHAP rank", "CNN rank", "Mean rank"], top_rows(cons, cons.columns, 15), widths=[1.4, 1.4, 1.4, 1.4])
    doc.add_picture(str(FIG / "fig_consensus_top.png"), width=Inches(6.2))
    add_caption(doc, "图 3. mRNA PAM50 共识关键基因 top-20")
    add_desc(doc, "共识排序综合 SHAP 和 CNN saliency 排名，排名靠前的基因更可能在两种模型解释中都具有重要性。")

    doc.add_heading("5. Survival 关键因子", level=1)
    sur = pd.read_csv(DATA / "survival_key_factors.tsv", sep="\t")
    sur_clean = sur.dropna(subset=["hr", "p"]).sort_values("p").head(10)
    add_table(doc, ["Gene", "LR abs coef", "Hazard ratio", "p-value"], top_rows(sur_clean, sur_clean.columns, 10), widths=[1.5, 1.5, 1.5, 1.4])
    doc.add_picture(str(FIG / "fig_survival_hr.png"), width=Inches(6.2))
    add_caption(doc, "图 4. Survival 单因素 Cox top-20 基因风险比")
    add_desc(doc, "风险比小于 1 的基因提示高表达与较好生存相关，大于 1 提示高风险。")

    doc.add_heading("6. 通路富集", level=1)
    enrich = pd.read_csv(DATA / "mrna_pam50_top100_enrichment.tsv", sep="\t")
    enrich_sub = enrich[["source", "term_id", "term_name", "p_value", "intersection_size"]]
    add_table(doc, ["Source", "Term ID", "Term name", "p-value", "Intersection size"], top_rows(enrich_sub, enrich_sub.columns, 10), widths=[0.8, 1.0, 2.5, 0.9, 0.9])
    doc.add_picture(str(FIG / "fig_enrichment.png"), width=Inches(6.2))
    add_caption(doc, "图 5. top-100 共识基因通路富集")
    add_desc(doc, "主要富集通路包括雌激素信号通路、细胞增殖调控、细胞分裂、上皮发育等，与乳腺癌分子分型具有生物学一致性。")

    doc.add_heading("7. JSD 图像化的创新点与可视化优势", level=1)
    doc.add_paragraph("FullSizeCNN 使用 JSD 排序后的灰度图，使基因重要性可以直接映射回图像中的像素位置。相比普通表格特征，JSD 图像能够同时展示原始表达、重要性排序、功能类别和模型关注区域。")
    doc.add_picture(str(FIG / "jsd_feature_mining_visualization.png"), width=Inches(6.2))
    add_caption(doc, "图 6. JSD 图像化特征挖掘可视化")
    add_desc(doc, "图中第一行展示灰度表达图、JSD 重要性图和基因功能类别图；第二行展示 SHAP 重要性图、CNN saliency 图和表达+saliency 叠加图。该可视化方式使得模型解释能够直接定位到具体基因像素，支持后续关键因子挖掘。")

    doc.add_heading("8. 结论", level=1)
    for c in [
        "SHAP 和 CNN saliency 均识别出 ESR1、TFF1、AGR3、FOXC1、KCNJ3 等乳腺癌相关基因。",
        "共识关键因子在 5 折中具有较好稳定性。",
        "Survival 单因素 Cox 发现 MS4A1、COL17A1、FABP7、CCL19 等与预后显著相关。",
        "通路富集结果支持关键基因与乳腺癌增殖、分化和激素信号相关。",
    ]:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(c)

    doc.save(OUT)
    print("Saved ->", OUT)


if __name__ == "__main__":
    main()

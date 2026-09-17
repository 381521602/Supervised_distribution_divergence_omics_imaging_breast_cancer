#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the imaging prediction summary DOCX."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
FIG = DATA / "images/report_figures"
OUT = ROOT / "组学图像化预测结果总结.docx"


def read(path):
    return pd.read_csv(DATA / path, sep="\t")


def fmt(row):
    return f"{row['mean']:.4f} ± {row['std']:.4f}"


def set_style(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5), 18, 10),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5), 14, 7),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78), 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)
    p.paragraph_format.space_after = Pt(10)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    return p


def add_description(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p.paragraph_format.space_after = Pt(8)
    return p


def set_table_width(table, widths_inches):
    table.autofit = False
    table.alignment = 1
    tbl = table._tbl
    tblPr = tbl.tblPr
    # remove old tblW if any
    for el in tblPr.findall(qn("w:tblW")):
        tblPr.remove(el)
    tblW = tblPr.makeelement(qn("w:tblW"), {qn("w:w"): str(int(sum(widths_inches) * 1440)), qn("w:type"): "dxa"})
    tblPr.append(tblW)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_inches):
                cell.width = Inches(widths_inches[idx])


def add_table(doc, headers, rows, widths=None, header_fill=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    if widths is None:
        widths = [6.5 / len(headers)] * len(headers)
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(9)
        if header_fill:
            shd = hdr[i]._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {qn("w:fill"): header_fill})
            hdr[i]._tc.get_or_add_tcPr().append(shd)
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9)
    return table


def table_from_df(df, omics_col=None, variant_col=None, task=None, metrics=None, variants=None, omics_list=None):
    rows = []
    if task:
        df = df[df["task"] == task]
    if omics_list:
        df = df[df[omics_col].isin(omics_list)]
    if variants:
        df = df[df[variant_col].isin(variants)]
    if omics_col:
        for o in omics_list or sorted(df[omics_col].unique()):
            row = [o]
            for m in metrics:
                sub = df[(df[omics_col] == o) & (df["metric"] == m)]
                if sub.empty:
                    row.append("")
                else:
                    row.append(fmt(sub.iloc[0]))
            rows.append(row)
    else:
        for v in variants or sorted(df[variant_col].unique()):
            row = [v]
            for m in metrics:
                sub = df[(df[variant_col] == v) & (df["metric"] == m)]
                row.append(fmt(sub.iloc[0]) if not sub.empty else "")
            rows.append(row)
    return rows


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    set_style(doc)

    add_title(doc, "基于组学图像化的预测结果总结")
    doc.add_paragraph("本文档汇总基于 JSD 图像化的 FullSizeCNN 及相关结构在 PAM50 分子分型和总生存期预测中的结果。所有结果均为基于预选特征的结果，采用 5 折交叉验证，报告均值 ± 标准差。")

    # 2 data setup
    doc.add_heading("1. 数据与图像化设置", level=1)
    doc.add_paragraph("各组学最终特征经 JSD 排序后，按中心向外的螺旋方式生成单通道灰度图像。图像尺寸如下：")
    add_table(
        doc,
        ["任务", "mRNA", "CNV", "miRNA"],
        [
            ["PAM50", "20×20", "8×8", "25×25"],
            ["Survival", "15×15", "13×13", "25×25"],
        ],
        widths=[1.2, 1.7, 1.7, 1.7],
        header_fill="E8EEF5",
    )
    add_description(doc, "表 1 列出了两个任务下三组学图像的边长。PAM50 任务中，mRNA 图像为 20×20，CNV 图像为 8×8，miRNA 图像为 25×25；Survival 任务中，mRNA 为 15×15，CNV 为 13×13，miRNA 为 25×25。不同组学图像尺寸不同，因此多组学整合时优先使用独立分支，而不是直接统一尺寸后堆叠。")

    # 3 FullSizeCNN architecture
    doc.add_heading("2. FullSizeCNN 网络结构与算法描述", level=1)
    doc.add_paragraph("FullSizeCNN 的核心是使用与输入图像尺寸相同的全尺寸卷积核，对整张单通道组学图像进行全局卷积。卷积后得到 32 个 1×1 特征图，再展平并输入小型 MLP 分类头。PAM50 任务输出 4 类概率，Survival 任务输出 1 个死亡风险分数。")
    doc.add_picture(str(FIG / "fullsize_cnn_architecture.png"), width=Inches(6.2))
    add_caption(doc, "图 1. FullSizeCNN 网络结构")
    add_description(doc, "图 1 展示了 FullSizeCNN 的主要计算流程：输入图像 H×W×1，经 H×W 全尺寸卷积核生成 32 个 1×1 特征图；随后展平为 32 维向量，并经过 64 维隐藏层、ReLU 激活和 Dropout，最后输出类别概率或生存风险分数。该结构参数少、与图像尺寸直接相关，适合当前样本量有限的组学图像化任务。")

    # 4 single-omics
    doc.add_heading("3. 单组学基础 FullSizeCNN", level=1)
    df_wgan = read("wgan_gp_augmentation_all_omics_results.tsv")
    base = df_wgan[df_wgan["variant"] == "baseline"]
    rows = table_from_df(base, omics_col="omics", omics_list=["mRNA", "CNV", "miRNA"], metrics=["accuracy", "macro_f1"])
    doc.add_heading("3.1 PAM50 分型", level=2)
    add_table(doc, ["组学", "Accuracy", "Macro-F1"], rows, widths=[1.4, 2.4, 2.7], header_fill="E8EEF5")
    add_description(doc, "表 2 显示，在 PAM50 分型中 mRNA 的判别能力最强，Accuracy 为 0.9316，Macro-F1 为 0.9205；miRNA 次之，Accuracy 为 0.8230，Macro-F1 为 0.7785；CNV 最弱，Accuracy 为 0.7164，Macro-F1 为 0.6727。该结果说明 mRNA 图像包含最稳定的 PAM50 分类信息。")
    doc.add_picture(str(FIG / "fig_single_omics_baseline.png"), width=Inches(6.2))
    add_caption(doc, "图 2. 单组学 FullSizeCNN 基线结果")
    add_description(doc, "图 2 左图展示 PAM50 分型结果，横轴为 mRNA、CNV、miRNA，纵轴为性能值。蓝色柱表示 Accuracy，浅蓝色柱表示 Macro-F1。mRNA 的 Accuracy 和 Macro-F1 均明显高于 CNV 和 miRNA。右图展示生存预测结果，深蓝色柱表示 ROC AUC，浅蓝色柱表示 C-index；mRNA 同样最高，CNV 和 miRNA 明显较低。该图直观反映 mRNA 在两类任务中的单组学优势。")

    rows = table_from_df(base, omics_col="omics", omics_list=["mRNA", "CNV", "miRNA"], metrics=["roc_auc", "c_index"])
    doc.add_heading("3.2 Survival 生存预测", level=2)
    add_table(doc, ["组学", "ROC AUC", "C-index"], rows, widths=[1.4, 2.4, 2.7], header_fill="E8EEF5")
    add_description(doc, "表 3 显示，Survival 任务中 mRNA 的 ROC AUC 为 0.6768、C-index 为 0.7125，均为三组学最高；CNV 和 miRNA 的判别能力较弱。该结果表明，图像化 FullSizeCNN 在生存预测中最主要的信号来自 mRNA。")

    # 5 classic structures
    doc.add_heading("4. 经典结构与轻量网络对比", level=1)
    df_cls = read("jsd_classic_structures_results.tsv")
    models = ["FullSizeCNN", "GroupNorm_Mish_FullSizeCNN", "FCN_FullSize", "ResNet_FullSize", "MultiBranch_SE_FullSize"]
    rows = table_from_df(df_cls, variant_col="variant", variants=models, metrics=["accuracy", "macro_f1"])
    doc.add_heading("4.1 PAM50 分型", level=2)
    add_table(doc, ["结构", "Accuracy", "Macro-F1"], rows, widths=[2.2, 2.1, 2.2], header_fill="E8EEF5")
    add_description(doc, "表 4 显示，基础 FullSizeCNN 在 PAM50 分型中仍最优，Accuracy 为 0.9316、Macro-F1 为 0.9205。GroupNorm+Mish、FCN、ResNet 轻量残差和多分支 SE 结构均未超过该基线。多分支 SE 结构表现最弱，说明当前小样本条件下增加过多分支和注意力模块反而可能降低稳定性。")
    rows = table_from_df(df_cls, variant_col="variant", variants=models, metrics=["roc_auc", "c_index"])
    doc.add_heading("4.2 Survival 生存预测", level=2)
    add_table(doc, ["结构", "ROC AUC", "C-index"], rows, widths=[2.2, 2.1, 2.2], header_fill="E8EEF5")
    add_description(doc, "表 5 显示，Survival 任务中 ResNet 轻量残差结构最优，ROC AUC 为 0.6897、C-index 为 0.7250。基础 FullSizeCNN 的 C-index 为 0.7125，略低于 ResNet；FCN 结构表现最弱。这提示在生存预测中加入轻量残差跳连有助于保留更有效的全局特征。")
    doc.add_picture(str(FIG / "fig_classic_structures.png"), width=Inches(6.2))
    add_caption(doc, "图 3. 经典全尺寸卷积结构对比")
    add_description(doc, "图 3 左图展示 PAM50 分型结果，横轴为 Base、GN+Mish、FCN、ResNet、SE 五种结构，蓝色柱为 Accuracy，浅蓝色柱为 Macro-F1。Base 的 Accuracy 最高，其他结构均略低。右图展示生存预测结果，深蓝色柱为 ROC AUC，浅蓝色柱为 C-index；ResNet 的 C-index 最高，FCN 的 C-index 最低。整体来看，结构复杂化并未在 PAM50 中带来收益，但轻量残差在生存预测中有一定优势。")

    doc.add_paragraph("轻量经典 CNN 中，PAM50 仍以基础 FullSizeCNN 最好；Survival 中 ResNet 轻量残差结构和 ShuffleNetV2Lite 较接近。DenseNetLite 与 LightInception 在当前小图、小样本下不稳定。")

    # 5 reorder
    doc.add_heading("5. 不同图像排序方案", level=1)
    df_reorder = read("reorder_fullsize_cnn_results.tsv")
    rows = table_from_df(df_reorder, variant_col="variant", metrics=["accuracy", "macro_f1"])
    doc.add_heading("5.1 PAM50 分型", level=2)
    add_table(doc, ["排序方案", "Accuracy", "Macro-F1"], rows, widths=[2.4, 2.0, 2.1], header_fill="E8EEF5")
    add_description(doc, "表 6 显示，原始 JSD 螺旋排序在 Accuracy 上最高，为 0.9316；方案 2 的 Macro-F1 最高，为 0.9243，但 Accuracy 略低。方案 1 和方案 2 与原始排序接近，方案 3 表现最弱。综合看，功能重排并未稳定超过 JSD 排序。")
    df_reorder_sur = read("reorder_fullsize_cnn_survival_results.tsv")
    rows = table_from_df(df_reorder_sur, variant_col="variant", metrics=["roc_auc", "c_index"])
    doc.add_heading("5.2 Survival 生存预测", level=2)
    add_table(doc, ["排序方案", "ROC AUC", "C-index"], rows, widths=[2.4, 2.0, 2.1], header_fill="E8EEF5")
    add_description(doc, "表 7 显示，Survival 任务中原始 JSD 螺旋排序仍最优，ROC AUC 为 0.6768、C-index 为 0.7125。方案 3 的 C-index 为 0.7066，最接近原始排序；方案 2 的 C-index 最低，为 0.6838。总体而言，重新规划基因空间顺序没有在生存预测中带来提升。")
    doc.add_paragraph("配对检验显示，各排序方案之间总体上未形成稳健显著差异，原始 JSD 螺旋排序仍是最稳默认方案。")

    # 6 augmentation
    doc.add_heading("6. GAN / WGAN-GP 数据扩增", level=1)
    doc.add_paragraph("扩增均在每折训练集内完成，测试集不参与生成和标准化。")
    doc.add_picture(str(FIG / "fig_augmentation.png"), width=Inches(6.2))
    add_caption(doc, "图 4. WGAN-GP 扩增对三组学预测的影响")
    add_description(doc, "图 4 左图展示 PAM50 分型中不同扩增倍数对 Accuracy 的影响，横轴为 None、1×、5×、10×，纵轴为 Accuracy。mRNA 在 5× 时达到最高点，CNV 和 miRNA 在 1× 时较高，之后随扩增倍数增加而下降。右图展示 Survival 任务中 C-index 的变化，miRNA 在 10× 时略有提高，CNV 随扩增倍数增加持续下降。该图说明扩增效果高度依赖组学和任务，不能统一设定扩增倍数。")
    doc.add_paragraph("PAM50 中，mRNA 在 5× 扩增时 Accuracy 较高；CNV、miRNA 在 1× 时相对最好。Survival 中，miRNA 的 ROC AUC 随扩增倍数略有改善，CNV 扩增总体无益。")

    # 7 integration
    doc.add_heading("7. 多组学 FullSizeCNN 整合", level=1)
    df_int = read("fullsize_multimodal_integration_results.tsv")
    schemes = ["single_mRNA", "single_CNV", "single_miRNA", "triple_concat", "triple_gated", "mrna_mirna_concat", "late_fusion_avg"]
    labels = ["mRNA", "CNV", "miRNA", "Triple concat", "Triple gated", "mRNA+miRNA", "Late fusion"]
    for i, s in enumerate(schemes):
        df_int.loc[df_int["scheme"] == s, "display"] = labels[i]
    rows = table_from_df(df_int, variant_col="display", variants=labels, metrics=["accuracy", "macro_f1"])
    doc.add_heading("7.1 PAM50 分型", level=2)
    add_table(doc, ["方案", "Accuracy", "Macro-F1"], rows, widths=[1.7, 2.3, 2.5], header_fill="E8EEF5")
    add_description(doc, "表 8 显示，在共同样本条件下，PAM50 分型中 mRNA 单组学 Accuracy 最高，为 0.9128；三分支门控融合的 Macro-F1 最高，为 0.8955。三分支拼接、mRNA+miRNA 和晚期概率平均融合均未超过门控融合。门控融合在 Accuracy 略低于 mRNA 单组学的情况下，改善了类别间平衡能力。")
    rows = table_from_df(df_int, variant_col="display", variants=labels, metrics=["roc_auc", "c_index"])
    doc.add_heading("7.2 Survival 生存预测", level=2)
    add_table(doc, ["方案", "ROC AUC", "C-index"], rows, widths=[1.7, 2.3, 2.5], header_fill="E8EEF5")
    add_description(doc, "表 9 显示，Survival 任务中单组学概率晚期平均融合最优，ROC AUC 为 0.7591、C-index 为 0.7503。mRNA+miRNA 两分支的 C-index 为 0.7415，是次优方案。三分支拼接和门控融合低于晚期平均融合，说明在生存预测中，先分别训练单组学模型再融合概率，比直接联合训练分支更稳定。")
    doc.add_picture(str(FIG / "fig_integration.png"), width=Inches(6.2))
    add_caption(doc, "图 5. FullSizeCNN 多组学整合结果")
    add_description(doc, "图 5 左图展示 PAM50 共同样本整合结果，蓝色柱为 Accuracy，浅蓝色柱为 Macro-F1。mRNA 单组学 Accuracy 最高，门控融合的 Macro-F1 最高。右图展示 Survival 共同样本整合结果，深蓝色柱为 ROC AUC，浅蓝色柱为 C-index。晚期平均融合的 ROC AUC 和 C-index 均最高，是当前生存预测中最优的图像化多组学方案。")
    doc.add_paragraph("PAM50 最优整合为三分支门控融合；Survival 最优整合为单组学概率晚期平均融合。")

    # 8 conclusion
    doc.add_heading("8. 结论", level=1)
    conclusions = [
        "PAM50 分型更适合直接使用 mRNA 单组学 FullSizeCNN，Accuracy 最高为 0.9316。",
        "Survival 生存预测更适合多组学概率晚期融合，ROC AUC 0.7591、C-index 0.7503 为当前图像化最优。",
        "原始 JSD 螺旋排序仍是最稳的图像化方式。",
        "WGAN-GP 扩增收益有限且依赖任务和组学，不建议作为统一默认策略。",
        "图像化 FullSizeCNN 的主要价值体现在生存预测和多组学融合。",
    ]
    for c in conclusions:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(c)

    doc.add_heading("9. 关键结果文件", level=1)
    for f in [
        "fullsize_multimodal_integration_results.tsv",
        "wgan_gp_augmentation_all_omics_results.tsv",
        "jsd_classic_structures_results.tsv",
        "reorder_fullsize_cnn_results.tsv",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"data/{f}")

    doc.save(OUT)
    print(f"Saved -> {OUT}")


if __name__ == "__main__":
    main()
